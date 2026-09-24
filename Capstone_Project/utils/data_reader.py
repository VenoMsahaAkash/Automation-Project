"""
data_reader.py
Utility for reading test data from both JSON and Excel (.xlsx) files.
"""

import json
import logging
import os

import openpyxl

from config.config import JSON_DATA_FILE, EXCEL_DATA_FILE

logger = logging.getLogger(__name__)


# ─── JSON Reader ──────────────────────────────────────────────────────────────

def read_json(filepath: str = JSON_DATA_FILE) -> dict:
    """
    Read and parse a JSON test data file.

    :param filepath: Path to the .json file (defaults to configured path).
    :return:         Parsed dictionary.
    :raises FileNotFoundError: If the file does not exist.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"JSON data file not found: {filepath}")

    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)

    logger.info(f"[DataReader] JSON data loaded from: {filepath}")
    return data


def get_json_value(key_path: str, filepath: str = JSON_DATA_FILE):
    """
    Get a nested value from the JSON data file using dot notation.

    :param key_path: Dot-separated key path (e.g., 'login.email').
    :param filepath: Path to the JSON file.
    :return:         The value at the given path.
    :raises KeyError: If the key path does not exist.

    Example:
        email = get_json_value('login.email')
    """
    data = read_json(filepath)
    keys = key_path.split(".")
    value = data
    for key in keys:
        value = value[key]
    logger.info(f"[DataReader] JSON key '{key_path}' = {value!r}")
    return value


# ─── Excel Reader ─────────────────────────────────────────────────────────────

def read_excel(sheet_name: str, filepath: str = EXCEL_DATA_FILE) -> list[dict]:
    """
    Read data from a specific Excel sheet into a list of row dictionaries.
    The first row is treated as the header.

    :param sheet_name: Name of the worksheet to read.
    :param filepath:   Path to the .xlsx file.
    :return:           List of dicts, one per data row.
    :raises FileNotFoundError: If the file does not exist.
    :raises KeyError: If the sheet name is not found.

    Example:
        rows = read_excel('Login')
        # [{'Field': 'email', 'Value': '...', 'Description': '...'}, ...]
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(
            f"Excel data file not found: {filepath}\n"
            f"Run 'python create_excel_data.py' to generate it."
        )

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)

    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")

    ws      = wb[sheet_name]
    rows    = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    data    = [dict(zip(headers, row)) for row in rows[1:]]

    logger.info(f"[DataReader] Excel sheet '{sheet_name}' loaded: {len(data)} rows.")
    wb.close()
    return data


def get_excel_value(sheet_name: str, field_name: str,
                    filepath: str = EXCEL_DATA_FILE) -> str:
    """
    Retrieve a value from the Excel sheet by matching the 'Field' column.
    Assumes the sheet has 'Field' and 'Value' columns.

    :param sheet_name:  Worksheet name (e.g., 'Login').
    :param field_name:  The value in the 'Field' column to look up.
    :param filepath:    Path to the .xlsx file.
    :return:            Corresponding 'Value' cell content as a string.
    :raises ValueError: If field_name is not found in the sheet.
    """
    rows = read_excel(sheet_name, filepath)
    for row in rows:
        if str(row.get("Field", "")).strip().lower() == field_name.strip().lower():
            value = str(row.get("Value", "")).strip()
            logger.info(f"[DataReader] Excel [{sheet_name}][{field_name}] = {value!r}")
            return value
    raise ValueError(f"Field '{field_name}' not found in sheet '{sheet_name}'.")
